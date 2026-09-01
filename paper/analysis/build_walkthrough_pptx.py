#!/usr/bin/env python3
"""Rebuild the 25-slide public/private-returns walkthrough deck from the current
analysis (2026-09-01 Arrhenius + pH-6.0 basis, net-export accounting).

The original deck was produced with PptxGenJS and its generator was not kept;
this script replaces it with a reproducible python-pptx builder so the deck can
be regenerated whenever the pipeline re-runs. Content decisions follow
paper/deck-update-checklist.md (carbon slides re-led on uniform-20; price, MRV
and delivered cost presented as a single (p-m)/c lever; allocation tags on all
results slides).

Inputs : docs/maps/envelopes/*.png (fresh erw-11/12 maps),
         paper/figures/*.png (overliming, pH profile),
         paper/figures/walkthrough_assets/*.png (reused illustrative assets
         extracted from the 2026-07 deck: schematic, framework, cost map,
         kinetic saturation, optimal rates),
         numbers hard-coded below from docs/tables/*.csv (2026-09-01).
Output : paper/ERW_public_private_returns_walkthrough.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
FIG = os.path.join(ROOT, "paper", "figures")
AST = os.path.join(FIG, "walkthrough_assets")
ENV = os.path.join(ROOT, "docs", "maps", "envelopes")
OUT = os.path.join(ROOT, "paper", "ERW_public_private_returns_walkthrough.pptx")

# ---- palette / type ---------------------------------------------------------
INK    = RGBColor(0x21, 0x26, 0x2B)
MUTED  = RGBColor(0x5A, 0x60, 0x67)
ACCENT = RGBColor(0x2E, 0x6B, 0x46)   # deep leaf green (private)
BLUE   = RGBColor(0x1A, 0x4F, 0x9C)   # carbon / public
RUST   = RGBColor(0xA0, 0x3D, 0x26)   # combined-only / emphasis
PAPER  = RGBColor(0xFB, 0xFA, 0xF7)
CARD   = RGBColor(0xF1, 0xEF, 0xE9)
LINE   = RGBColor(0xD9, 0xD5, 0xCC)
FONT   = "Avenir Next"
FONT_FB = "Calibri"

SW, SH = Inches(13.333), Inches(7.5)
FOOTER = "Enhanced Rock Weathering — public & private returns in SSA    ·    Gebrekidan & Chamberlin (2026)"

prs = Presentation()
prs.slide_width, prs.slide_height = SW, SH
BLANK = prs.slide_layouts[6]
_n = [0]

def slide(bg=PAPER):
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    _n[0] += 1
    return s

def _set(run, size, color=INK, bold=False, italic=False):
    f = run.font
    f.name, f.size, f.bold, f.italic = FONT, Pt(size), bold, italic
    f.color.rgb = color

def box(s, x, y, w, h):
    tb = s.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tb, tf

def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, space_after=4, anchor=None):
    """runs: list of paragraphs; each paragraph a list of (txt, size, color, bold, italic)."""
    tb, tf = box(s, x, y, w, h)
    if anchor: tf.vertical_anchor = anchor
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space_after)
        for (t, size, color, bold, italic) in para:
            r = p.add_run(); r.text = t
            _set(r, size, color, bold, italic)
    return tb

def eyebrow(s, txt, color=ACCENT, x=Inches(0.55), y=Inches(0.32)):
    text(s, x, y, Inches(12), Inches(0.32),
         [[(txt.upper(), 12.5, color, True, False)]])

def title(s, txt, y=Inches(0.62), size=30, color=INK, w=Inches(12.2)):
    text(s, Inches(0.52), y, w, Inches(0.9), [[(txt, size, color, True, False)]])

def footer(s):
    text(s, Inches(0.55), Inches(7.08), Inches(11.2), Inches(0.3),
         [[(FOOTER, 9, MUTED, False, False)]])
    text(s, Inches(12.35), Inches(7.08), Inches(0.6), Inches(0.3),
         [[(str(_n[0]), 10, MUTED, True, False)]], align=PP_ALIGN.RIGHT)

def tag(s, txt, color=BLUE):
    """allocation/regime tag, top right"""
    tb = text(s, Inches(8.6), Inches(0.34), Inches(4.2), Inches(0.3),
              [[(txt.upper(), 10.5, color, True, False)]], align=PP_ALIGN.RIGHT)
    return tb

def card(s, x, y, w, h, fill=CARD):
    from pptx.enum.shapes import MSO_SHAPE
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sh.adjustments[0] = 0.055
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = LINE; sh.line.width = Pt(0.75)
    sh.shadow.inherit = False
    return sh

def pic(s, path, x, y, max_w, max_h, align="center"):
    iw, ih = Image.open(path).size
    scale = min(max_w / iw, max_h / ih)
    w, h = int(iw * scale), int(ih * scale)
    px = x + int((max_w - w) / 2) if align == "center" else x
    py = y + int((max_h - h) / 2)
    return s.shapes.add_picture(path, px, py, w, h)

def stat(s, x, y, w, big, unit, label, color=ACCENT, big_size=40):
    text(s, x, y, w, Inches(1.15),
         [[(big, big_size, color, True, False), (" " + unit, 16, color, True, False)],
          [(label, 12.5, MUTED, False, False)]], space_after=2)

def bullets(s, x, y, w, h, items, head_size=14.5, body_size=12.5, gap=10):
    runs = []
    for head, body in items:
        runs.append([(head, head_size, INK, True, False)])
        runs.append([(body, body_size, MUTED, False, False)])
    tb = text(s, x, y, w, h, runs, space_after=4)
    for i, p in enumerate(tb.text_frame.paragraphs):
        if i % 2 == 1: p.space_after = Pt(gap)
    return tb

def divider(num, big, sub):
    s = slide(bg=RGBColor(0x1E, 0x2A, 0x24))
    text(s, Inches(0.9), Inches(2.35), Inches(3.0), Inches(0.5),
         [[(num.upper(), 15, RGBColor(0x8F, 0xBA, 0x84), True, False)]])
    text(s, Inches(0.9), Inches(2.85), Inches(11.5), Inches(1.1),
         [[(big, 40, RGBColor(0xF4, 0xF2, 0xEC), True, False)]])
    text(s, Inches(0.9), Inches(4.05), Inches(10.8), Inches(1.4),
         [[(sub, 15, RGBColor(0xB9, 0xC2, 0xBB), False, False)]])
    text(s, Inches(12.35), Inches(7.08), Inches(0.6), Inches(0.3),
         [[(str(_n[0]), 10, RGBColor(0x8A, 0x93, 0x8C), True, False)]], align=PP_ALIGN.RIGHT)

# ============================================================ 1 · title
s = slide()
text(s, Inches(0.9), Inches(1.15), Inches(11.5), Inches(0.4),
     [[("EX ANTE ANALYSIS  ·  FULL WALKTHROUGH  ·  2026-09 MODEL BASIS", 13, ACCENT, True, False)]])
text(s, Inches(0.9), Inches(1.75), Inches(11.6), Inches(1.9),
     [[("Mapping the public and private returns to", 34, INK, True, False)],
      [("Enhanced Rock Weathering in Sub-Saharan Africa", 34, INK, True, False)]], space_after=2)
text(s, Inches(0.9), Inches(3.75), Inches(11.5), Inches(0.5),
     [[("Bisrat Haile Gebrekidan", 16, INK, True, False), ("   (CIMMYT-Ethiopia)      ·      ", 13, MUTED, False, False),
       ("Jordan Chamberlin", 16, INK, True, False), ("   (CIMMYT-Kenya)", 13, MUTED, False, False)]])
text(s, Inches(0.9), Inches(4.55), Inches(11.5), Inches(0.4),
     [[("A spatially explicit, decomposed profitability model  —  44 SSA countries  ·  23 crops  ·  per-pixel grid", 14, MUTED, False, False)]])
text(s, Inches(0.9), Inches(5.9), Inches(11.5), Inches(0.8),
     [[("Working draft for Agricultural Economics — coauthor review.  All numbers from the 2026-09-01 run:", 12, MUTED, False, False)],
      [("Arrhenius weathering kinetics (Ea = 68.8 kJ/mol)  ·  pH optimum 6.0  ·  net-export CDR accounting  ·  ESROC-verified sources", 12, ACCENT, False, False)]], space_after=2)

# ============================================================ 2 · proposition
s = slide(); eyebrow(s, "The proposition")
title(s, "ERW pays the farmer twice")
text(s, Inches(0.55), Inches(1.28), Inches(5.6), Inches(0.75),
     [[("Spreading finely ground basalt on cropland couples two value streams that are usually separate:", 13, MUTED, False, False)]])
bullets(s, Inches(0.55), Inches(2.1), Inches(5.6), Inches(2.6), [
    ("PRIVATE — agronomic", "Alkalinity neutralizes soil acidity like agricultural lime, raising yields. Accrues to the farmer."),
    ("PUBLIC — carbon (CDR)", "Weathering durably removes atmospheric CO₂, monetizable through voluntary carbon markets (VCM)."),
])
text(s, Inches(0.55), Inches(5.0), Inches(5.6), Inches(1.6),
     [[("Central question:  ", 13.5, INK, True, False),
       ("where does each return justify deployment — and where does neither suffice alone, yet together they do?", 13.5, MUTED, False, False)]])
pic(s, os.path.join(AST, "image-2-1.png"), Inches(6.4), Inches(1.3), Inches(6.5), Inches(5.4))
footer(s)

# ============================================================ 3 · questions
s = slide(); eyebrow(s, "What we ask")
title(s, "Six research questions")
QS = [
    ("Q1", "Where does the private (yield) return alone justify ERW investment?"),
    ("Q2", "Where does the public (carbon) return alone justify it?"),
    ("Q3", "Where do the two coincide — the intersection?"),
    ("Q4", "Where does neither stream alone suffice, but the combined margin is positive?"),
    ("Q5", "How robust is the deployable geography to parameter assumptions?"),
    ("Q6", "What would shift the public (carbon) envelope outward?"),
]
for i, (q, t) in enumerate(QS):
    col, row = i % 2, i // 2
    x = Inches(0.55) + col * Inches(6.35)
    y = Inches(1.55) + row * Inches(1.7)
    card(s, x, y, Inches(6.05), Inches(1.45))
    text(s, x + Inches(0.25), y + Inches(0.18), Inches(0.9), Inches(1.0),
         [[(q, 24, ACCENT, True, False)]])
    text(s, x + Inches(1.1), y + Inches(0.16), Inches(4.75), Inches(1.15),
         [[(t, 13, INK, False, False)]], anchor=MSO_ANCHOR.MIDDLE)
footer(s)

# ============================================================ 4 · framework
s = slide(); eyebrow(s, "Modeling framework")
title(s, "A decomposed per-pixel gross margin")
text(s, Inches(0.55), Inches(1.35), Inches(5.3), Inches(0.8),
     [[("π  =  R", 24, INK, True, False), ("agro", 14, ACCENT, True, False),
       ("  +  R", 24, INK, True, False), ("CDR", 14, BLUE, True, False),
       ("  -  C", 24, INK, True, False)],
      [("gross margin = private agronomic return + public carbon return - full delivered cost", 11.5, MUTED, False, True)]],
     space_after=2)
bullets(s, Inches(0.55), Inches(2.55), Inches(5.3), Inches(4.2), [
    ("Resolution", "~0.083° grid × 44 SSA countries × 23 SPAM crops, area-weighted to each pixel."),
    ("Three streams", "Agronomic yield uplift, durable CDR credit, and a fully decomposed delivered cost."),
    ("Decision object", "Each pixel is classified by which stream(s) make the margin positive — not a single supply curve."),
    ("Regimes", "Evaluated at year-1, NPV @ 10%, and equilibrium steady state (the headline)."),
], body_size=12, gap=8)
pic(s, os.path.join(AST, "image-4-1.png"), Inches(6.05), Inches(1.25), Inches(7.0), Inches(5.6))
footer(s)

# ============================================================ 5 · inputs
s = slide(); eyebrow(s, "Inputs")
title(s, "Public data layers behind every pixel")
LAYERS = [
    ("SOIL", "SoilGrids", "pH, ECEC, texture on cropland"),
    ("CROPS", "SPAM v2", "23 crop distributions & production"),
    ("AGRONOMY", "EcoCrop", "crop pH-suitability response"),
    ("FEEDSTOCK", "GLiM", "basalt source geology"),
    ("LOGISTICS", "MAP friction 2019", "travel-time cost surface (Weiss et al. 2020)"),
    ("CLIMATE", "WorldClim", "temperature & moisture for weathering"),
    ("PRICES", "FAOSTAT", "2016–20 producer-price medians"),
    ("LIME REQ.", "LiTAS", "lime-equivalent dose (targeting)"),
    ("ACIDITY SINK", "Merlos / Kamprath", "standing & maintenance acidity for net-export"),
]
for i, (k, src, d) in enumerate(LAYERS):
    col, row = i % 3, i // 3
    x = Inches(0.55) + col * Inches(4.25)
    y = Inches(1.55) + row * Inches(1.72)
    card(s, x, y, Inches(4.0), Inches(1.5))
    text(s, x + Inches(0.22), y + Inches(0.14), Inches(3.6), Inches(1.25),
         [[(k, 10.5, ACCENT, True, False)],
          [(src, 14.5, INK, True, False)],
          [(d, 11, MUTED, False, False)]], space_after=1)
footer(s)

# ============================================================ 6 · divider I
divider("Part I", "The assumption sets",
        "Five modeling choices govern every result: delivered cost, carbon price, net-export CDR "
        "accounting, an over-liming private penalty, and kinetic weathering saturation.")

# ============================================================ 7 · cost lever
s = slide(); eyebrow(s, "Assumption 1 · Delivered cost")
title(s, "Cost is dominated by transport — a real lever")
bullets(s, Inches(0.55), Inches(1.5), Inches(5.4), Inches(4.6), [
    ("Decomposed cost", "Quarry-gate + grinding + freight + spreading, built up per pixel from the friction surface; MRV netted from the carbon price."),
    ("Transport binds", "Mean 55% (median 66%) of delivered cost. Remote acid soils are priced out even where the agronomy is strong."),
    ("The solar lever", "Electrified / solar-assisted haulage cuts the variable transport term: -25% transport widens the private envelope 4.31 → 5.12 Mha, the public one only 1.95 → 2.07."),
], gap=10)
pic(s, os.path.join(AST, "image-7-1.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 8 · cost table
s = slide(); eyebrow(s, "Assumption 1 · Delivered cost")
title(s, "Delivered basalt cost, decomposed")
ROWS = [
    ("Component", "Cost per t basalt", "Type", "Basis / source", True),
    ("Quarry-gate feedstock", "$10 / t", "fixed", "basalt fines as aggregate by-product (cf. $17.50/t US crushed-stone average, USGS)", False),
    ("Grinding (50 µm)", "19.4 kWh/t × local tariff (≈$1.9/t)", "fixed*", "Strefler et al. 2018 energy curve", False),
    ("Transport (haulage)", "travel-time × $0.04/min/t", "variable", "30 t truck over MAP friction surface; within SSA corridor tariffs (Teravaninthorn & Raballand 2009)", False),
    ("Spreading", "$8 / t", "fixed", "on-farm application", False),
    ("Delivered cost / t", "quarry + grinding + transport + spreading", "= sum", "the delivered $/t consumed per pixel", True),
    ("MRV (netted separately)", "$20 / tCO₂", "per tCO₂", "low end of the $15–71/tCO₂ supplier range (Mercer et al. 2024, LSE Grantham)", False),
]
y = Inches(1.5)
for comp, cost, typ, basis, hdr in ROWS:
    if hdr:
        card(s, Inches(0.55), y, Inches(12.25), Inches(0.62), fill=RGBColor(0xE4, 0xE9, 0xE2))
    text(s, Inches(0.75), y + Inches(0.09), Inches(2.6), Inches(0.5),
         [[(comp, 12, INK, True, False)]])
    text(s, Inches(3.45), y + Inches(0.09), Inches(3.35), Inches(0.5),
         [[(cost, 11.5, INK, hdr, False)]])
    text(s, Inches(6.9), y + Inches(0.09), Inches(0.95), Inches(0.5),
         [[(typ, 11, MUTED, False, False)]])
    text(s, Inches(7.9), y + Inches(0.05), Inches(4.85), Inches(0.62),
         [[(basis, 10, MUTED, False, False)]])
    y += Inches(0.72)
text(s, Inches(0.55), y + Inches(0.15), Inches(12.2), Inches(0.5),
     [[("* grinding varies with the country electricity tariff; its life-cycle CO₂ (with transport and spreading) is separately deducted from gross removal.", 10.5, MUTED, False, True)]])
footer(s)

# ============================================================ 9 · carbon price
s = slide(); eyebrow(s, "Assumption 2 · Carbon price")
title(s, "The carbon return uses a net VCM price")
tag(s, "equilibrium · uniform-20 (carbon lead)")
text(s, Inches(0.55), Inches(1.45), Inches(5.4), Inches(1.0),
     [[("R", 20, INK, True, False), ("CDR", 12, BLUE, True, False),
       ("  =  durable CDR × ( p - m )", 20, INK, True, False)],
      [("p = VCM carbon price ($/tCO₂)  ·  m = per-tonne MRV cost", 11.5, MUTED, False, True)]], space_after=2)
text(s, Inches(0.55), Inches(2.7), Inches(5.4), Inches(2.2),
     [[("Baseline is ", 12.5, MUTED, False, False), ("$150/tCO₂", 12.5, INK, True, False),
       (" — a conservative floor: 2024–25 durable-CDR transactions cluster at $250–450/tCO₂ (CDR.fyi). "
        "Price, MRV and delivered cost move the public envelope only through the single ratio (p - m)/c.", 12.5, MUTED, False, False)],
      [("The ladder is steeply convex: between $250 and $350 the deployable area jumps from 7 to 23 Mha.", 12.5, INK, False, False)]],
     space_after=8)
LAD = [("$/tCO₂", "Public Mha", "% treated", "Durable CDR (Mt)", True),
       ("$100", "1.40", "4.3%", "7.8", False),
       ("$150", "3.95", "12.1%", "20.4", False),
       ("$250", "7.00", "21.5%", "33.3", False),
       ("$350", "23.06", "70.7%", "100.6", False),
       ("$500", "27.98", "85.8%", "115.6", False)]
y = Inches(1.5)
for a, b, c, d, hdr in LAD:
    if hdr:
        card(s, Inches(6.5), y, Inches(6.3), Inches(0.56), fill=RGBColor(0xE0, 0xE7, 0xF0))
    for xi, txt_, wd in [(6.7, a, 1.2), (8.0, b, 1.5), (9.6, c, 1.3), (11.0, d, 1.7)]:
        text(s, Inches(xi), y + Inches(0.08), Inches(wd), Inches(0.45),
             [[(txt_, 12.5, BLUE if hdr else INK, hdr, False)]])
    y += Inches(0.66)
text(s, Inches(6.5), y + Inches(0.12), Inches(6.3), Inches(0.9),
     [[("Treated base: 32.6 Mha under uniform-20. On the targeted dose the $150 envelope is 1.95 Mha (11.1% of its 17.6 Mha treated base) — shares are not comparable across allocations.",
        10.5, MUTED, False, True)]])
footer(s)

# ============================================================ 10 · net-export
s = slide(); eyebrow(s, "Assumption 3 · Net-export CDR")
title(s, "Credit only carbon that is durably exported")
text(s, Inches(0.55), Inches(1.42), Inches(5.35), Inches(1.0),
     [[("durable CDR = max( gross - F·S , 0 )", 19, INK, True, False)],
      [("F = 0.88 tCO₂ per t CaCO₃-eq  ·  S = acidity sink (standing / maintenance)", 11.5, MUTED, False, True)]], space_after=2)
text(s, Inches(0.55), Inches(2.5), Inches(5.35), Inches(1.35),
     [[("Alkalinity spent neutralizing soil acidity re-releases its CO₂ — exactly as agricultural lime does — so it must be deducted "
        "(West & McBride 2005; Hamilton et al. 2007; Holden et al. 2024). A first-order sequential bound; exchange-stored cations can export later (Kanzaki et al. 2025).",
        11.5, MUTED, False, False)]])
stat(s, Inches(0.55), Inches(4.0), Inches(2.6), "≈70%", "", "of credited removal is cut at first application once acidity neutralization is netted out", color=RUST, big_size=36)
text(s, Inches(3.4), Inches(4.0), Inches(2.5), Inches(1.6),
     [[("Durable removal is largely a ", 12, MUTED, False, False), ("steady-state phenomenon", 12, INK, True, False),
       (" — equilibrium retention is 77% (targeted) to 90% (uniform-20). Hence the equilibrium headline.", 12, MUTED, False, False)]])
pic(s, os.path.join(FIG, "walkthrough_ph_profile_npv.png"), Inches(6.15), Inches(1.35), Inches(6.9), Inches(5.35))
footer(s)

# ============================================================ 11 · over-liming
s = slide(); eyebrow(s, "Assumption 4 · Over-liming penalty")
title(s, "Too much alkalinity can depress yield")
bullets(s, Inches(0.55), Inches(1.5), Inches(5.4), Inches(4.9), [
    ("Mechanism", "Beyond the lime requirement, pH is pushed past the crop optimum; the EcoCrop response applies a yield haircut at the post-application pH."),
    ("Buffered shift", "pH change is damped by soil buffering (β = kβ·ECEC); reacted alkalinity A = gross CDR / F."),
    ("Effect", "Concentrated on high uniform doses: the central estimate trims the targeted private return ~7% (band 6–10%), but 20% at uniform-50 with a 3–99% band."),
    ("Status", "The penalty parameters are stated assumptions — the ERW literature documents the risk only qualitatively (Swoboda et al. 2022)."),
], gap=8)
pic(s, os.path.join(FIG, "overliming_sensitivity.png"), Inches(6.15), Inches(1.35), Inches(6.9), Inches(5.4))
footer(s)

# ============================================================ 12 · saturation
s = slide(); eyebrow(s, "Assumption 5 · Kinetic saturation")
title(s, "Weathering reacts less-than-linearly with dose")
text(s, Inches(0.55), Inches(1.45), Inches(5.4), Inches(1.0),
     [[("sat(D) = ( K", 19, INK, True, False), ("d", 12, INK, True, False),
       (" + 50 ) / ( K", 19, INK, True, False), ("d", 12, INK, True, False), (" + D )", 19, INK, True, False)],
      [("effective reacted tonnage saturates in applied dose D, normalized at the Lewis (2021) 50 t/ha reference", 11.5, MUTED, False, True)]], space_after=2)
text(s, Inches(0.55), Inches(2.75), Inches(5.4), Inches(3.4),
     [[("Key finding.  ", 13, INK, True, False),
       ("Saturation makes CDR concave in dose (giving an interior optimum), but it does not drive the public envelope to zero. "
        "Below the 50 t/ha reference, per-tonne reactivity is higher — so at targeted doses durable CDR actually rises. "
        "What binds the carbon case is MAC > price, not kinetics.", 13, MUTED, False, False)]])
pic(s, os.path.join(AST, "image-12-1.png"), Inches(6.15), Inches(1.35), Inches(6.9), Inches(5.4))
footer(s)

# ============================================================ 13 · design space
s = slide(); eyebrow(s, "Design space")
title(s, "Temporal regimes × allocation rules")
text(s, Inches(0.55), Inches(1.4), Inches(12.2), Inches(0.6),
     [[("The headline is the ", 13, MUTED, False, False), ("equilibrium steady state with net-export accounting", 13, INK, True, False),
       (" — the targeted dose leads the private and intersection geography; ", 13, MUTED, False, False),
       ("uniform-20 leads the carbon case", 13, BLUE, True, False),
       (" (the lime-requirement dose is CDR-minimal by construction).", 13, MUTED, False, False)]])
text(s, Inches(0.55), Inches(2.25), Inches(6.0), Inches(0.35), [[("TEMPORAL REGIMES", 12, ACCENT, True, False)]])
bullets(s, Inches(0.55), Inches(2.6), Inches(6.0), Inches(3.6), [
    ("Year-1", "One-time dose & cost; durable CDR net of standing acidity. Most conservative on carbon."),
    ("NPV @ 10%", "Discounted stream of relief and durable removal, phased over the 5-year residency."),
    ("Equilibrium", "Steady-state durable removal against the maintenance dose — the headline regime."),
], gap=7)
text(s, Inches(6.9), Inches(2.25), Inches(6.0), Inches(0.35), [[("ALLOCATION RULES", 12, BLUE, True, False)]])
bullets(s, Inches(6.9), Inches(2.6), Inches(6.0), Inches(3.6), [
    ("Targeted", "LiTAS lime-equivalent dose — agronomically matched. Leads private & intersection."),
    ("Uniform 10 / 20 / 50 t/ha", "Fixed rates — trace how CDR and over-liming scale with dose. Uniform-20 leads the carbon case."),
    ("Weathering kinetics", "Arrhenius temperature response (Ea = 68.8 kJ/mol) · unimodal pH factor peaking at 6.0 (2026-09 basis)."),
], gap=7)
footer(s)

# ============================================================ 14 · establishment vs maintenance
s = slide(); eyebrow(s, "How to read the returns")
title(s, "Establishment dose vs. steady-state maintenance")
bars = [("Year 1\n(establishment)", 1.0, RUST), ("Year 2", 0.18, MUTED), ("Year 3", 0.18, MUTED),
        ("Year 4", 0.18, MUTED), ("steady state\n(maintenance)", 0.18, ACCENT)]
bx = Inches(0.9)
for lbl, hfrac, col in bars:
    h = Emu(int(Inches(2.6) * hfrac))
    card(s, bx, Inches(1.7) + (Inches(2.6) - h), Inches(1.65), h, fill=col)
    text(s, bx - Inches(0.1), Inches(4.45), Inches(1.9), Inches(0.7),
         [[(lbl, 10, MUTED, False, False)]], align=PP_ALIGN.CENTER)
    bx += Inches(2.05)
text(s, Inches(0.9), Inches(1.32), Inches(11.0), Inches(0.35),
     [[("bar height ∝ basalt applied that period — charged at its full delivered cost; the yield benefit recurs every year once soils are held at target",
        10.5, MUTED, False, True)]])
bullets(s, Inches(0.9), Inches(5.15), Inches(11.6), Inches(1.8), [
    ("EQUILIBRIUM — the headline", "Full annual yield vs the fully-costed maintenance top-up, once soils are at target. Private envelope: ≈ $3.5 B yield for ≈ $1.74 B rock (~2:1 per cycle). Excludes the establishment dose."),
    ("YEAR-1 — establishment", "One-time full lime-equivalent dose; standing acidity consumes most first-application alkalinity, so the carbon case is thin until steady state."),
], gap=6)
footer(s)

# ============================================================ 15 · divider II
divider("Part II", "Results",
        "Envelopes and the combined-only band · durable carbon and marginal abatement cost · "
        "what shifts the public frontier · optimal rates · the robust core geography.")

# ============================================================ 16 · typology
s = slide(); eyebrow(s, "Results · Typology")
title(s, "A five-way pixel classification")
tag(s, "equilibrium · net-export · targeted")
TYP = [("Private-sufficient", "R_agro - C > 0", ACCENT),
       ("Public-sufficient", "R_CDR - C > 0", BLUE),
       ("Intersection", "both hold at the same pixel", RGBColor(0x6A, 0x3D, 0x8F)),
       ("Combined-only", "neither alone, but GM > 0", RUST),
       ("Neither", "GM ≤ 0 — not deployable", MUTED)]
y = Inches(1.55)
for name, cond, col in TYP:
    card(s, Inches(0.55), y, Inches(5.0), Inches(0.92))
    text(s, Inches(0.8), y + Inches(0.12), Inches(2.6), Inches(0.7),
         [[(name, 13.5, col, True, False)]], anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(3.15), y + Inches(0.12), Inches(2.3), Inches(0.7),
         [[(cond, 11, MUTED, False, False)]], anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(1.04)
pic(s, os.path.join(ENV, "env_typology_targeted.png"), Inches(5.9), Inches(1.3), Inches(7.15), Inches(5.5))
footer(s)

# ============================================================ 17 · private envelope
s = slide(); eyebrow(s, "Results · Private envelope")
title(s, "Where yield alone pays")
tag(s, "equilibrium · targeted", ACCENT)
text(s, Inches(0.55), Inches(1.45), Inches(5.3), Inches(0.9),
     [[("Land where the agronomic return covers the full delivered cost with no carbon credit at all.", 13, MUTED, False, False)]])
stat(s, Inches(0.55), Inches(2.5), Inches(2.6), "4.3", "Mha", "private-sufficient area", color=ACCENT)
stat(s, Inches(3.3), Inches(2.5), Inches(2.6), "$3.5", "B", "private return per cycle", color=ACCENT)
text(s, Inches(0.55), Inches(4.1), Inches(5.3), Inches(2.2),
     [[("The largest single envelope under the targeted dose — and unchanged by the carbon-side model revision, because it does not depend on the CDR surface.",
        12.5, MUTED, False, False)],
      [("But it is a minority of acid cropland: most viable land needs more than yield to pencil.", 12.5, INK, False, False)]], space_after=8)
pic(s, os.path.join(ENV, "env_private_targeted.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 18 · public envelope
s = slide(); eyebrow(s, "Results · Public envelope")
title(s, "Where carbon alone pays (durably)")
tag(s, "equilibrium · uniform-20 (carbon lead)")
text(s, Inches(0.55), Inches(1.45), Inches(5.3), Inches(0.9),
     [[("Land where durable, net-export CDR at $150/tCO₂ covers the full delivered cost with no yield benefit.", 13, MUTED, False, False)]])
stat(s, Inches(0.55), Inches(2.5), Inches(2.6), "3.95", "Mha", "durable public area (20.4 Mt)", color=BLUE)
stat(s, Inches(3.3), Inches(2.5), Inches(2.6), "12.1%", "", "of treated cropland", color=BLUE)
text(s, Inches(0.55), Inches(4.1), Inches(5.3), Inches(2.3),
     [[("On the targeted dose the carbon-only case is 1.95 Mha (11.1% of its treated base) — the lime-requirement dose is CDR-minimal by construction.",
        12.5, MUTED, False, False)],
      [("Net-export accounting and today's price keep the carbon-only case tightly bound — the reason the combined case matters.", 12.5, INK, False, False)]], space_after=8)
pic(s, os.path.join(ENV, "env_public_uniform_20.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 19 · combined-only
s = slide(); eyebrow(s, "Results · The central result", color=RUST)
title(s, "The combined-only band is the policy object")
tag(s, "equilibrium · targeted", ACCENT)
stat(s, Inches(0.55), Inches(1.5), Inches(3.2), "2.7", "Mha", "combined-only", color=RUST, big_size=44)
text(s, Inches(0.55), Inches(2.75), Inches(5.3), Inches(0.95),
     [[("neither the yield return nor the carbon return alone covers cost — yet together the margin is positive.", 13, MUTED, False, False)]])
for i, (v, lbl) in enumerate([("4.3", "private (Mha)"), ("2.0", "public (Mha)"), ("7.6", "combined (Mha)")]):
    stat(s, Inches(0.55) + i * Inches(1.85), Inches(3.9), Inches(1.7), v, "", lbl, color=INK, big_size=26)
text(s, Inches(0.55), Inches(5.15), Inches(5.3), Inches(1.6),
     [[("The economic point:  ", 12.5, INK, True, False),
       ("voluntary carbon finance is not a bonus on an already-profitable practice — it is the marginal revenue that unlocks an otherwise-unprofitable, yield-improving investment.",
        12.5, MUTED, False, False)]])
pic(s, os.path.join(ENV, "env_intersection_targeted.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 20 · MAC
s = slide(); eyebrow(s, "Results · What binds the carbon case")
title(s, "Marginal abatement cost, not kinetics")
tag(s, "equilibrium · uniform-20 (carbon lead)")
stat(s, Inches(0.55), Inches(1.5), Inches(4.2), "$276", "/ tCO₂", "area-weighted median MAC — vs a $130 net credit price ($330 on the targeted dose)", color=BLUE, big_size=42)
text(s, Inches(0.55), Inches(3.15), Inches(5.3), Inches(1.3),
     [[("MAC = delivered cost per tonne of rock ÷ durable CDR per tonne. The public envelope is exactly the land where MAC < (p - m).",
        12.5, MUTED, False, False)],
      [("That median sits at the upper edge of the published cost-synthesis IQR ($137–276; Suhrhoff et al. 2026) — the residual gap is accounting stringency, not SSA logistics.",
        12.5, INK, False, False)]], space_after=8)
text(s, Inches(0.55), Inches(5.1), Inches(5.3), Inches(1.5),
     [[("So the carbon case is limited by ", 12.5, MUTED, False, False),
       ("how much CO₂ a tonne of rock verifiably removes", 12.5, INK, True, False),
       (" — the largest and least-resolved lever, and the clear field-calibration priority.", 12.5, MUTED, False, False)]])
pic(s, os.path.join(ENV, "mac_uniform_20.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 21 · two levers
s = slide(); eyebrow(s, "Results · The levers")
title(s, "Two levers — not three — shift the public envelope")
tag(s, "equilibrium · uniform-20 (carbon lead)")
bullets(s, Inches(0.55), Inches(1.5), Inches(5.4), Inches(4.9), [
    ("1 · The price–cost ratio (p - m)/c",
     "Carbon price, MRV and delivered cost enter the public test only through this single ratio: deployable area climbs from 4.3% of treated cropland at $100 to 86% at $500. Cost cuts (solar haulage) act on the same lever — and move mostly the private envelope."),
    ("2 · Durable CDR per tonne of rock",
     "Scaled where it physically acts — on gross removal, before the acidity sink. Because the sink and life-cycle terms are fixed subtrahends, durable removal is super-linear in it: doubling the rate takes the public envelope 3.95 → 22.5 Mha (~6×). Kinetic saturation is second-order here."),
], gap=12)
pic(s, os.path.join(ENV, "cdr_rate_sensitivity_panel_uniform_20.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 22 · optimal rates
s = slide(); eyebrow(s, "Results · Optimal application rates")
title(s, "Each objective has an interior optimum")
text(s, Inches(0.55), Inches(1.5), Inches(5.4), Inches(4.4),
     [[("With over-liming and saturation in place, each per-pixel objective is concave — so a pixel-specific ", 13, MUTED, False, False),
       ("optimal rate", 13, INK, True, False),
       (" exists for the private, public, and combined margins.", 13, MUTED, False, False)],
      [("The combined optimum is moderate (typically well below blanket-high uniform doses): enough over-application to earn extra carbon, but not so much that over-liming and saturation erode the return.",
        13, MUTED, False, False)],
      [("Prototype refinement — maps are illustrative (pre-revision basis); promoting them to the headline is on the roadmap.",
        11, MUTED, False, True)]], space_after=10)
pic(s, os.path.join(AST, "image-22-1.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 23 · robust core
s = slide(); eyebrow(s, "Results · Robust core")
title(s, "A defensible priority geography")
tag(s, "equilibrium · net-export · targeted", ACCENT)
text(s, Inches(0.55), Inches(1.45), Inches(5.3), Inches(0.85),
     [[("Intersecting the envelope across the 180-point parameter grid isolates land that stays viable under most assumptions.", 12.5, MUTED, False, False)]])
stat(s, Inches(0.55), Inches(2.45), Inches(2.5), "0.65", "Mha", "robust core (≥50% of grid)", color=RUST, big_size=32)
stat(s, Inches(3.15), Inches(2.45), Inches(2.5), "0.38", "M", "farms reached", color=RUST, big_size=32)
stat(s, Inches(0.55), Inches(3.85), Inches(2.5), "$0.60", "B", "crop value at stake", color=RUST, big_size=32)
stat(s, Inches(3.15), Inches(3.85), Inches(2.5), "1.22", "Mha", "candidate tier (≥33%)", color=RGBColor(0xC0, 0x7A, 0x2B), big_size=32)
text(s, Inches(0.55), Inches(5.3), Inches(5.3), Inches(1.4),
     [[("Four-fifths of the candidate tier sits in Cameroon (0.61 Mha) and Guinea (0.38) — the geography to deploy first, and the base to expand as the price–cost ratio or the verified CDR rate improves.",
        12.5, MUTED, False, False)]])
pic(s, os.path.join(ENV, "core_priority.png"), Inches(6.15), Inches(1.3), Inches(6.9), Inches(5.5))
footer(s)

# ============================================================ 24 · takeaways
s = slide(); eyebrow(s, "Takeaways")
title(s, "ERW is a coupled agronomic-and-carbon proposition")
TK = [
    ("1 · The two returns rarely each suffice alone",
     "But together they make a meaningful area of SSA cropland profitable that neither would — the 2.7 Mha combined-only band."),
    ("2 · Carbon finance is the margin, not a bonus",
     "On most viable land, VCM revenue is what unlocks a yield-improving investment. That reframes ERW away from a pure CDR supply curve."),
    ("3 · The binding constraint is durable CDR per tonne",
     "Net-export accounting cuts first-application removal ~70%; the area-weighted median MAC is $276 (uniform-20) / $330 (targeted). The price–cost ratio and the verified CDR rate are the only two structural levers."),
    ("4 · Treating it as coupled changes the answer",
     "Both where one would deploy first (the Cameroon and Guinea highlands core) and what one would invest in to expand it (field-calibrated removal rates and cheaper delivery)."),
]
for i, (h, b) in enumerate(TK):
    col, row = i % 2, i // 2
    x = Inches(0.55) + col * Inches(6.35)
    y = Inches(1.6) + row * Inches(2.55)
    card(s, x, y, Inches(6.05), Inches(2.3))
    text(s, x + Inches(0.28), y + Inches(0.2), Inches(5.5), Inches(1.95),
         [[(h, 14, INK, True, False)],
          [(b, 11.5, MUTED, False, False)]], space_after=5)
footer(s)

# ============================================================ 25 · caveats / next
s = slide(); eyebrow(s, "Next")
title(s, "Caveats and where we go next")
text(s, Inches(0.55), Inches(1.5), Inches(6.0), Inches(0.35), [[("CAVEATS", 12, RUST, True, False)]])
bullets(s, Inches(0.55), Inches(1.85), Inches(6.0), Inches(4.6), [
    ("Priorities", "Field calibration of durable removal per tonne is the single highest-value input — it sets MAC and the whole public case."),
    ("Not yet modeled", "Tenure and upfront finance; compaction; downstream river–ocean losses (~15–20% of exported removal, Renforth & Henderson 2017); country-level credit eligibility (Article 6). Omitted N₂O effects likely run in ERW's favor."),
    ("Farm counts", "Based on national average holding sizes — order-of-magnitude."),
], gap=8)
text(s, Inches(6.9), Inches(1.5), Inches(6.0), Inches(0.35), [[("NEXT", 12, ACCENT, True, False)]])
bullets(s, Inches(6.9), Inches(1.85), Inches(6.0), Inches(4.6), [
    ("Registry-grade accounting", "Align the net-export bound with VCM methodology counterfactuals (liming substitution, financial-additionality screens)."),
    ("Optimal-rate maps", "Promote the pixel-specific private/public/combined optima from prototype to headline."),
    ("Share & collaborate", "The interactive envelope explorer reproduces every slider of this analysis in one self-contained HTML file — no install."),
], gap=8)
footer(s)

prs.save(OUT)
print(f"wrote {OUT} ({_n[0]} slides)")

# document properties (avoid python-pptx template defaults)
prs2 = Presentation(OUT)
cp = prs2.core_properties
cp.title = "Public and private returns to Enhanced Rock Weathering in Sub-Saharan Africa"
cp.author = "Bisrat Haile Gebrekidan; Jordan Chamberlin"
cp.last_modified_by = "Bisrat Haile Gebrekidan"
cp.comments = "Full walkthrough deck, regenerated from the 2026-09-01 model run by build_walkthrough_pptx.py"
prs2.save(OUT)
print("core properties set")
